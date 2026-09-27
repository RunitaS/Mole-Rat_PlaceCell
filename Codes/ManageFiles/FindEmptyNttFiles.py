import os
import re
import sys

from openpyxl import Workbook

# ---- Settings ----
ROOT_DIR = r"X:\NMR_group_data\Runita\Data\Ephys_Data\AllSortedData\Tetrode"   # <-- change this
SAVE_XLSX = True                               # also write results as a 3-sheet .xlsx into ROOT_DIR

# Neuralynx .ntt layout: 16384-byte (16 KB) text header, then 304-byte spike records.
# A file with zero spikes is therefore exactly the header size.
HEADER_BYTES = 16384
RECORD_BYTES = 304


TT_PATTERN = re.compile(r"TT(\d+)", re.IGNORECASE)   # tetrode id at the start of the filename


def tt_sort_key(tt):
    m = TT_PATTERN.match(tt)
    return (0, int(m.group(1))) if m else (1, tt)


def n_spikes(fpath):
    return max(0, (os.path.getsize(fpath) - HEADER_BYTES) // RECORD_BYTES)


def find_empty_ntt(root_dir):
    if not os.path.isdir(root_dir):
        sys.exit(f"Not a valid directory: {root_dir}")

    empty_files = []          # every .ntt with zero spikes
    all_empty_folders = []    # folders where every .ntt is empty
    empty_tetrodes = []       # (folder, tetrode, n_files): every file of that tetrode is empty,
                              # but the folder still has spikes on other tetrodes
    n_checked, n_failed = 0, 0

    for dirpath, dirnames, filenames in os.walk(root_dir):
        ntt_files = [f for f in filenames if f.lower().endswith(".ntt")]
        if not ntt_files:
            continue

        n_empty_here = 0
        tt_counts = {}        # tetrode -> [n_files, n_empty]
        for fname in ntt_files:
            fpath = os.path.join(dirpath, fname)
            m = TT_PATTERN.match(fname)
            tt = f"TT{m.group(1)}" if m else os.path.splitext(fname)[0]
            counts = tt_counts.setdefault(tt, [0, 0])
            try:
                counts[0] += 1
                if n_spikes(fpath) == 0:
                    empty_files.append(fpath)
                    n_empty_here += 1
                    counts[1] += 1
                n_checked += 1
            except Exception as e:
                print(f"FAILED: {fpath} -> {e}")
                n_failed += 1

        if n_empty_here == len(ntt_files):
            all_empty_folders.append((dirpath, len(ntt_files)))
        else:
            for tt in sorted(tt_counts, key=tt_sort_key):
                n_files, n_empty = tt_counts[tt]
                if n_files > 0 and n_empty == n_files:
                    empty_tetrodes.append((dirpath, tt, n_files))

    return empty_files, all_empty_folders, empty_tetrodes, n_checked, n_failed


if __name__ == "__main__":
    empty_files, all_empty_folders, empty_tetrodes, n_checked, n_failed = find_empty_ntt(ROOT_DIR)

    print(f"\n=== Empty .ntt files ({len(empty_files)}) ===")
    for f in empty_files:
        print(f)

    print(f"\n=== Folders where ALL .ntt files are empty ({len(all_empty_folders)}) ===")
    for d, n in all_empty_folders:
        print(f"{d}   ({n} .ntt files)")

    print(f"\n=== Tetrodes with ALL .ntt files empty (other tetrodes in folder have spikes) "
          f"({len(empty_tetrodes)}) ===")
    for d, tt, n in empty_tetrodes:
        print(f"{d}   {tt}   ({n} .ntt files)")

    print(f"\nChecked {n_checked} .ntt files. Empty: {len(empty_files)}. "
          f"Fully corrupted folders: {len(all_empty_folders)}. "
          f"Fully empty tetrodes: {len(empty_tetrodes)}. Failed to read: {n_failed}.")

    if SAVE_XLSX:
        out_xlsx = os.path.join(ROOT_DIR, "Empty_ntt_report.xlsx")
        wb = Workbook()

        ws_files = wb.active
        ws_files.title = "Empty_ntt_files"
        ws_files.append(["Folder", "FileName", "FullPath"])
        for f in empty_files:
            ws_files.append([os.path.dirname(f), os.path.basename(f), f])

        ws_folders = wb.create_sheet("AllEmpty_ntt_folders")
        ws_folders.append(["Folder", "N_ntt_files"])
        for d, n in all_empty_folders:
            ws_folders.append([d, n])

        ws_tetrodes = wb.create_sheet("AllEmpty_tetrodes")
        ws_tetrodes.append(["Folder", "Tetrode", "N_ntt_files"])
        for d, tt, n in empty_tetrodes:
            ws_tetrodes.append([d, tt, n])

        wb.save(out_xlsx)
        print(f"Saved: {out_xlsx}")
