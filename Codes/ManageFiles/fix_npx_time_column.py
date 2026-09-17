"""
Goes through every .xlsx file in a hardcoded directory and, for every sheet,
checks whether cell A1 already holds "time" or 0/"0".

If neither is present, column A is assumed to be missing its header, so the
script:
    1. Shifts all of column A's data down by one cell.
    2. Writes the text "time" into A1.
    3. Drops the original last entry of column A (so the column length is
       unchanged instead of growing by one row).

Requires: openpyxl  (pip install openpyxl)
"""

import os
from openpyxl import load_workbook

# ---- hardcode your directory here ----
DIRECTORY = r"C:/Runita/NMR/analysis/TrackingCorrection"


def needs_time_header(value):
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip().lower() != "time" and value.strip() != "0"
    if isinstance(value, (int, float)):
        return value != 0
    return True


def fix_sheet(ws):
    a1 = ws["A1"].value

    if not needs_time_header(a1):
        return False

    max_row = ws.max_row
    col_a_values = [ws.cell(row=r, column=1).value for r in range(1, max_row + 1)]

    # Drop the last entry, then shift everything down by one row.
    shifted_values = col_a_values[:-1]

    ws.cell(row=1, column=1, value="time")
    for i, val in enumerate(shifted_values):
        ws.cell(row=i + 2, column=1, value=val)

    return True


def main():
    for filename in os.listdir(DIRECTORY):
        if not filename.lower().endswith(".xlsx"):
            continue

        filepath = os.path.join(DIRECTORY, filename)
        wb = load_workbook(filepath)

        modified = False
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if fix_sheet(ws):
                modified = True
                print(f"Fixed '{sheet_name}' in {filename}")

        if modified:
            wb.save(filepath)


if __name__ == "__main__":
    main()
