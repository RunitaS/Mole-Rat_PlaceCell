# -*- coding: utf-8 -*-
"""
Bursting-cell detection for every Neuralynx .ntt file under ROOT_FOLDER.

File selection and loading follow Fig1def_spike_quality_v2.py: every .ntt file
found by os.walk(ROOT_FOLDER) is one unit, its session is the folder path
relative to ROOT_FOLDER, and spikes in cluster 0 (cell_number == 0, unsorted /
noise) are excluded.

Burst detection follows BurstFreqRange_v4.py (KDE of the ISI distribution,
burst power in a burst ISI window, KDE peak + half-maximum bounds), with
these fixes:
  - Cluster 0 is excluded (v4 used every spike in the file).
  - Fixed KDE bandwidth in ms (KDE_BW_MS) instead of Scott's rule. Scott's rule
    on ISIs up to 200 ms gives a bandwidth of several to tens of ms, which
    smears the few-ms burst peak away.
  - Binned KDE (fine ISI histogram convolved with a Gaussian): the same
    estimate as gaussian_kde on a grid, but fast for units with 10^5 spikes.
  - The burst peak must be an interior maximum of the KDE inside the burst
    window. v4 took the window argmax as the peak even when it sat on the
    window edge (KDE still rising at the upper bound = mode outside the window).
  - burst_power_pct is counted directly from the ISIs, not from the KDE area.
  - MIN_ISIS raised from 5 to 50: with a narrow bandwidth the KDE of a handful
    of ISIs is a set of spikes, not a distribution.

Per unit
  1. ISIs (ms) of the sorted spike times; ISIs in (0, ISI_MAX_MS] are kept.
  2. burst_power_pct = % of kept ISIs in [ISI_BURST_LO_MS, ISI_BURST_HI_MS].
  3. KDE of the kept ISIs; peak = highest KDE point in the burst window, and it
     must not lie on the window edge.
  4. has_burst_peak = burst_power_pct >= SIGNIFICANCE_PCT and an interior peak.
  5. Half-maximum bounds around the peak (linear interpolation, clamped to the
     window) -> burst ISI range and frequency range (Hz = 1000 / ISI_ms).
  6. CSI (same definition as Fig1def_spike_quality_v2.py).
     is_bursting = has_burst_peak and CSI > CSI_THRESHOLD
     (CSI condition dropped if REQUIRE_COMPLEX_SPIKES is False).

poisson_burst_pct is the % of ISIs (<= ISI_MAX_MS) that a Poisson cell with the
unit's mean rate would have in the burst window. High-rate tonic cells pass
SIGNIFICANCE_PCT (and CSI) without bursting, so compare burst_power_pct with it.
It is only used in the classification if POISSON_EXCESS_RATIO is set.

Outputs (in OUTPUT_DIR)
  BURST_DETECTION.xlsx                 per-unit table + summary sheet
  burst_plots/bursting/                ISI histogram + KDE of each bursting unit
  burst_plots/not_bursting/            same for the other units (PLOT_ALL_UNITS)
  burst_plots/BURST_DETECTION_summary  CSI / burst power / peak ISI / peak frequency
"""

import os
import numpy as np # type: ignore
import pandas as pd # type: ignore
import matplotlib # type: ignore
matplotlib.use('Agg')
import matplotlib.pyplot as plt # type: ignore
from matplotlib.colors import to_hex # type: ignore
from openpyxl.styles import PatternFill, Font, Alignment # type: ignore
from openpyxl.utils import get_column_letter # type: ignore
from openpyxl.drawing.image import Image as XLImage # type: ignore

# ── Configuration ──────────────────────────────────────────────────────────────

# Same root folder as Fig1def_spike_quality_v2.py
ROOT_FOLDER  = r'X:\NMR_group_data\Runita\Analysis\Thesis\Corr_Data'
OUTPUT_DIR   = ROOT_FOLDER
OUTPUT_EXCEL = os.path.join(OUTPUT_DIR, 'BURST_DETECTION.xlsx')
PLOT_DIR     = os.path.join(OUTPUT_DIR, 'burst_plots')

EXCLUDE_CLUSTER0 = True     # drop cell_number == 0 spikes, as in Fig1def_spike_quality_v2.py

# True: ISI plot for every unit (2000+ files take a while). False: bursting units only.
PLOT_ALL_UNITS = False

# ── Complex spike filter (BurstFreqRange_v4.py) ───────────────────────────────
CSI_THRESHOLD          = 10.0   # CSI must exceed this (%) for a cell to count as complex
# True: is_bursting also requires CSI > CSI_THRESHOLD.
# False: is_bursting = has_burst_peak only (CSI is still computed and saved).
REQUIRE_COMPLEX_SPIKES = False

# ── Burst ISI window (BurstFreqRange_v4.py) ───────────────────────────────────
ISI_BURST_LO_MS  = 1.0     # lower ISI bound of burst window (ms)
ISI_BURST_HI_MS  = 30.0    # upper ISI bound of burst window (ms)
SIGNIFICANCE_PCT = 5.0     # min % of ISIs in the burst window to call it significant
ISI_MAX_MS       = 200.0   # ISIs above this are ignored (also the KDE grid ceiling)
MIN_ISIS         = 50      # min number of ISIs <= ISI_MAX_MS to run the burst test
# None: classification as in BurstFreqRange_v4.py. A number, e.g. 2.0: has_burst_peak
# also requires burst_power_pct >= POISSON_EXCESS_RATIO x poisson_burst_pct, which
# rejects high-rate tonic cells whose short ISIs are expected by chance.
POISSON_EXCESS_RATIO = None

# ── KDE ───────────────────────────────────────────────────────────────────────
KDE_BIN_MS = 0.05          # grid resolution (ms)
KDE_BW_MS  = 1.0           # Gaussian kernel SD (ms)

# ── Per-unit ISI plot ─────────────────────────────────────────────────────────
PLOT_BIN_MS  = 0.5
PLOT_XMAX_MS = 80.0
PLOT_BINS_MS = np.linspace(0.0, PLOT_XMAX_MS, int(round(PLOT_XMAX_MS / PLOT_BIN_MS)) + 1)

# ── Figure palette (magma colormap + neutral inks, as in Fig1def_spike_quality_v2.py) ──
FIG_CMAP = 'magma'
MAGMA_1, MAGMA_2, MAGMA_3, MAGMA_4 = [to_hex(matplotlib.colormaps[FIG_CMAP](v))
                                      for v in np.linspace(0, 1, 6)[1:-1]]
REJECT_COLOR   = '#e1e0d9'
INK_PRIMARY    = '#0b0b0b'
INK_SECONDARY  = '#52514e'
INK_MUTED      = '#898781'
GRID_COLOR     = '#e1e0d9'
SURFACE        = '#fcfcfb'

BURST_COLOR     = MAGMA_2
NON_BURST_COLOR = REJECT_COLOR
KDE_COLOR       = MAGMA_3
WINDOW_COLOR    = MAGMA_4

os.makedirs(PLOT_DIR, exist_ok=True)


def _style_axes(ax):
    ax.set_facecolor(SURFACE)
    for side in ('top', 'right'):
        ax.spines[side].set_visible(False)
    for side in ('left', 'bottom'):
        ax.spines[side].set_color(INK_MUTED)
    ax.tick_params(colors=INK_SECONDARY, labelsize=9)
    ax.yaxis.grid(True, color=GRID_COLOR, linewidth=0.8)
    ax.set_axisbelow(True)


def _save(fig, path_base, exts=('png', 'svg')):
    """Save fig as <path_base>.<ext> for each ext, then close it."""
    for ext in exts:
        fig.savefig(f'{path_base}.{ext}', dpi=300, bbox_inches='tight', facecolor=SURFACE)
    plt.close(fig)


def _legend_outside(ax):
    ax.legend(frameon=False, fontsize=8, labelcolor=INK_PRIMARY,
              loc='upper left', bbox_to_anchor=(1.0, 1.0))


# ── Neuralynx .ntt loading (as in Fig1def_spike_quality_v2.py) ────────────────

NTT_HEADER_BYTES   = 16 * 1024
DEFAULT_ADBITVOLTS = 0.000000195

NTT_DTYPE = np.dtype([
    ('timestamp',   '<u8'),
    ('sc_number',   '<u4'),
    ('cell_number', '<u4'),
    ('params',      '<u4', (8,)),
    ('waveforms',   '<i2', (32, 4)),
])


def _parse_adbitvolts(path: str) -> float:
    try:
        with open(path, 'rb') as fh:
            header_raw = fh.read(NTT_HEADER_BYTES)
        header_txt = header_raw.decode('latin-1', errors='replace')
        for line in header_txt.splitlines():
            if 'ADBitVolts' in line:
                for p in line.split():
                    try:
                        val = float(p)
                        if 0 < val < 1:
                            return val
                    except ValueError:
                        pass
    except Exception:
        pass
    return DEFAULT_ADBITVOLTS


def load_ntt_waveforms(ntt_path: str, exclude_cluster0: bool = EXCLUDE_CLUSTER0):
    """
    Waveforms (µV, (n, 32, 4)) and timestamps (µs) from an .ntt file.

    exclude_cluster0 : if True, spikes with cell_number == 0 (unclustered /
        noise cluster) are dropped.
    """
    uv_scale   = _parse_adbitvolts(ntt_path) * 1e6
    spike_data = np.memmap(ntt_path, dtype=NTT_DTYPE, mode='r', offset=NTT_HEADER_BYTES)
    if exclude_cluster0:
        spike_data = spike_data[spike_data['cell_number'] != 0]  # type: ignore[index]
    waveforms  = spike_data['waveforms'].astype(np.float64) * uv_scale  # type: ignore[operator]
    timestamps = spike_data['timestamp'].astype(np.float64)
    return waveforms, timestamps


# ── Complex Spike Index (identical to Fig1def_spike_quality_v2.py) ────────────

def compute_csi(waveforms: np.ndarray, spike_ts_us: np.ndarray,
                min_isi_us: float = 3000.0, max_isi_us: float = 20000.0) -> float:
    """
    Complex Spike Index: % of ALL consecutive ISIs that fall in [3, 20] ms
    AND whose second spike has a smaller amplitude than the first.
    (Bhatt et al. 2020 / Ranck 1973)
    """
    if len(spike_ts_us) < 2:
        return float('nan')
    sort_idx  = np.argsort(spike_ts_us)
    ts_sorted = spike_ts_us[sort_idx]
    wf_sorted = waveforms[sort_idx]

    isis     = np.diff(ts_sorted)
    in_range = (isis >= min_isi_us) & (isis <= max_isi_us)

    if not np.any(in_range):
        return 0.0

    mean_pp = (wf_sorted.max(axis=1) - wf_sorted.min(axis=1)).mean(axis=1)  # (n_spikes,)
    first_pp  = mean_pp[:-1][in_range]
    second_pp = mean_pp[1:][in_range]

    n_complex = int(np.sum(second_pp < first_pp))
    return round(float(n_complex / len(isis) * 100.0), 4)


# ── Burst detection ───────────────────────────────────────────────────────────

BURST_KEYS = ['n_isis', 'burst_power_pct', 'poisson_burst_pct', 'has_burst_peak',
              'peak_isi_ms', 'peak_freq_hz', 'lower_isi_ms', 'upper_isi_ms',
              'lower_freq_hz', 'upper_freq_hz']


def _empty_burst() -> dict:
    out = {k: float('nan') for k in BURST_KEYS}
    out['n_isis'] = 0
    out['has_burst_peak'] = False
    return out


def isi_kde(isis_ms: np.ndarray):
    """
    Binned Gaussian KDE of the ISIs on [0, ISI_MAX_MS].
    Returns (bin centres in ms, density in 1/ms).
    """
    n_bins  = int(round(ISI_MAX_MS / KDE_BIN_MS))
    edges   = np.linspace(0.0, ISI_MAX_MS, n_bins + 1)
    counts, _ = np.histogram(isis_ms, bins=edges)

    half    = int(np.ceil(4.0 * KDE_BW_MS / KDE_BIN_MS))
    k       = np.arange(-half, half + 1) * KDE_BIN_MS
    kernel  = np.exp(-0.5 * (k / KDE_BW_MS) ** 2)
    kernel /= kernel.sum()
    smooth  = np.convolve(counts, kernel, mode='same')

    centres = 0.5 * (edges[:-1] + edges[1:])
    return centres, smooth / (len(isis_ms) * KDE_BIN_MS)


def _half_max_bounds(x: np.ndarray, y: np.ndarray, p: int):
    """Half-maximum crossings left and right of peak index p (window edges if none)."""
    half  = y[p] / 2.0
    lower = float(x[0])
    upper = float(x[-1])
    left  = np.flatnonzero(y[:p] <= half)
    if left.size:
        i = left[-1]                       # y[i] <= half < y[i + 1]
        lower = float(np.interp(half, [y[i], y[i + 1]], [x[i], x[i + 1]]))
    right = np.flatnonzero(y[p + 1:] <= half)
    if right.size:
        j = p + 1 + right[0]               # y[j] <= half < y[j - 1]
        upper = float(np.interp(half, [y[j], y[j - 1]], [x[j], x[j - 1]]))
    return lower, upper


def compute_burst(spike_ts_us: np.ndarray,
                  lo_ms: float = ISI_BURST_LO_MS,
                  hi_ms: float = ISI_BURST_HI_MS,
                  sig_pct: float = SIGNIFICANCE_PCT):
    """
    Burst power, KDE peak and half-maximum range of the burst ISIs.
    Returns (metrics dict, (kde_x, kde_y) or None).
    """
    out = _empty_burst()
    if len(spike_ts_us) < 2:
        return out, None

    isis_all = np.diff(np.sort(spike_ts_us)) / 1000.0          # µs → ms
    isis_all = isis_all[isis_all > 0]
    isis     = isis_all[isis_all <= ISI_MAX_MS]
    out['n_isis'] = int(len(isis))
    if len(isis) < MIN_ISIS:
        return out, None

    burst_power = 100.0 * float(np.mean((isis >= lo_ms) & (isis <= hi_ms)))
    out['burst_power_pct'] = round(burst_power, 3)

    # Poisson reference with the same mean rate, conditioned on ISI <= ISI_MAX_MS
    rate = 1.0 / float(np.mean(isis_all))                       # spikes / ms
    out['poisson_burst_pct'] = round(
        100.0 * (np.exp(-rate * lo_ms) - np.exp(-rate * hi_ms)) / (1.0 - np.exp(-rate * ISI_MAX_MS)), 3)

    kde_x, kde_y = isi_kde(isis)
    win    = (kde_x >= lo_ms) & (kde_x <= hi_ms)
    wx, wy = kde_x[win], kde_y[win]
    p      = int(np.argmax(wy))
    interior_peak = 0 < p < len(wy) - 1
    above_poisson = (POISSON_EXCESS_RATIO is None or
                     burst_power >= POISSON_EXCESS_RATIO * out['poisson_burst_pct'])
    if burst_power < sig_pct or not interior_peak or not above_poisson:
        return out, (kde_x, kde_y)

    lower, upper = _half_max_bounds(wx, wy, p)
    lower = max(lower, lo_ms)
    upper = min(upper, hi_ms)
    peak  = float(wx[p])
    out.update({
        'has_burst_peak': True,
        'peak_isi_ms':    round(peak, 3),
        'peak_freq_hz':   round(1000.0 / peak, 2),
        'lower_isi_ms':   round(lower, 3),
        'upper_isi_ms':   round(upper, 3),
        'lower_freq_hz':  round(1000.0 / upper, 2),   # long ISI = low frequency
        'upper_freq_hz':  round(1000.0 / lower, 2),
    })
    return out, (kde_x, kde_y)


def analyse_unit(ntt_path: str):
    """All burst metrics of one .ntt file. Returns (metrics, isis_ms, kde)."""
    waveforms, ts = load_ntt_waveforms(ntt_path)
    n_spikes = len(ts)
    duration_s = (ts.max() - ts.min()) / 1e6 if n_spikes >= 2 else 0.0

    csi        = compute_csi(waveforms, ts)
    is_complex = bool(not np.isnan(csi) and csi > CSI_THRESHOLD)
    burst, kde = compute_burst(ts)

    metrics = {
        'n_spikes':       n_spikes,
        'firing_rate_hz': round(n_spikes / duration_s, 4) if duration_s > 0 else float('nan'),
        'csi':            csi,
        'is_complex':     is_complex,
        **burst,
        'is_bursting':    bool(burst['has_burst_peak'] and (is_complex or not REQUIRE_COMPLEX_SPIKES)),
    }
    isis_ms = np.diff(np.sort(ts)) / 1000.0 if n_spikes >= 2 else np.array([])
    return metrics, isis_ms, kde


# ── Per-unit ISI plot ─────────────────────────────────────────────────────────

def plot_unit_isi(isis_ms: np.ndarray, kde, m: dict, title: str, path_base: str):
    """ISI histogram with KDE, burst window, and peak / half-max lines."""
    fig, ax = plt.subplots(figsize=(8, 4.5), facecolor=SURFACE)
    _style_axes(ax)

    shown = isis_ms[(isis_ms > 0) & (isis_ms <= PLOT_XMAX_MS)]
    ax.hist(shown, bins=PLOT_BINS_MS, color=BURST_COLOR if m['is_bursting'] else INK_MUTED,
            edgecolor=SURFACE, linewidth=0.4, zorder=3, label='ISI histogram')
    if kde is not None:
        kde_x, kde_y = kde
        keep = kde_x <= PLOT_XMAX_MS
        # density (1/ms) → spikes per histogram bin
        ax.plot(kde_x[keep], kde_y[keep] * m['n_isis'] * PLOT_BIN_MS, color=KDE_COLOR,
                linewidth=1.8, zorder=5, label=f'KDE (bw = {KDE_BW_MS} ms)')

    ax.axvspan(ISI_BURST_LO_MS, ISI_BURST_HI_MS, color=WINDOW_COLOR, alpha=0.15, lw=0, zorder=1,
               label=f'Burst window [{ISI_BURST_LO_MS:g}–{ISI_BURST_HI_MS:g} ms]')
    if m['has_burst_peak']:
        ax.axvline(m['peak_isi_ms'], color=INK_PRIMARY, linewidth=1.5, zorder=6,
                   label=f'Peak {m["peak_isi_ms"]:.1f} ms ({m["peak_freq_hz"]:.1f} Hz)')
        ax.axvline(m['lower_isi_ms'], color=INK_PRIMARY, linewidth=1.2, linestyle=':', zorder=6,
                   label=f'Half max {m["lower_isi_ms"]:.1f}–{m["upper_isi_ms"]:.1f} ms '
                         f'({m["lower_freq_hz"]:.1f}–{m["upper_freq_hz"]:.1f} Hz)')
        ax.axvline(m['upper_isi_ms'], color=INK_PRIMARY, linewidth=1.2, linestyle=':', zorder=6)

    def _pct(v):
        return 'N/A' if v is None or np.isnan(v) else f'{v:.2f}%'

    note = (f'CSI = {_pct(m["csi"])}  (threshold {CSI_THRESHOLD:g}%)\n'
            f'Burst power = {_pct(m["burst_power_pct"])}  (threshold {SIGNIFICANCE_PCT:g}%)\n'
            f'Poisson expectation = {_pct(m["poisson_burst_pct"])}\n'
            f'Rate = {m["firing_rate_hz"]:.2f} Hz,  n spikes = {m["n_spikes"]}\n'
            f'{"BURSTING" if m["is_bursting"] else "not bursting"}')
    ax.text(1.02, 0.0, note, transform=ax.transAxes, fontsize=8, color=INK_SECONDARY,
            ha='left', va='bottom')

    ax.set_xlim(0, PLOT_XMAX_MS)
    ax.set_xlabel('ISI (ms)', color=INK_PRIMARY)
    ax.set_ylabel(f'Spikes per {PLOT_BIN_MS:g} ms bin', color=INK_PRIMARY)
    ax.set_title(title, color=INK_PRIMARY, fontsize=10, loc='left')
    _legend_outside(ax)
    _save(fig, path_base, exts=('png',))


# ── Scan every .ntt file under ROOT_FOLDER ────────────────────────────────────

burst_plot_dir     = os.path.join(PLOT_DIR, 'bursting')
non_burst_plot_dir = os.path.join(PLOT_DIR, 'not_bursting')
os.makedirs(burst_plot_dir, exist_ok=True)
if PLOT_ALL_UNITS:
    os.makedirs(non_burst_plot_dir, exist_ok=True)

records = []
for dirpath, _, filenames in os.walk(ROOT_FOLDER):
    ntt_files = sorted(f for f in filenames if f.lower().endswith('.ntt'))
    if not ntt_files:
        continue

    session_name = os.path.relpath(dirpath, ROOT_FOLDER)
    for ntt_file in ntt_files:
        ntt_path = os.path.join(dirpath, ntt_file)
        print(f'Processing: {session_name}  |  {ntt_file}')
        try:
            metrics, isis_ms, kde = analyse_unit(ntt_path)
        except Exception as exc:
            print(f'  ERROR: {exc}')
            metrics = {'n_spikes': 0, 'firing_rate_hz': float('nan'), 'csi': float('nan'),
                       'is_complex': False, **_empty_burst(), 'is_bursting': False}
            isis_ms, kde = np.array([]), None

        records.append({'session': session_name, 'unit': ntt_file, **metrics})
        if metrics['is_bursting']:
            print(f'  BURSTING: CSI = {metrics["csi"]:.2f}%, burst power = {metrics["burst_power_pct"]:.1f}%, '
                  f'peak {metrics["peak_isi_ms"]:.1f} ms ({metrics["peak_freq_hz"]:.1f} Hz)')

        if metrics['n_spikes'] >= 2 and (metrics['is_bursting'] or PLOT_ALL_UNITS):
            parts = [p for p in (session_name, os.path.splitext(ntt_file)[0]) if p != '.']
            label = '_'.join(parts).replace(os.sep, '_').replace('/', '_')
            out_dir = burst_plot_dir if metrics['is_bursting'] else non_burst_plot_dir
            plot_unit_isi(isis_ms, kde, metrics, f'{session_name} | {ntt_file}',
                          os.path.join(out_dir, f'{label}_ISI'))

COLUMN_ORDER = ['session', 'unit', 'n_spikes', 'firing_rate_hz', 'csi', 'is_complex',
                'n_isis', 'burst_power_pct', 'poisson_burst_pct', 'has_burst_peak', 'is_bursting',
                'peak_isi_ms', 'peak_freq_hz', 'lower_isi_ms', 'upper_isi_ms',
                'lower_freq_hz', 'upper_freq_hz']
df = pd.DataFrame(records, columns=COLUMN_ORDER)

# ── Summary figure ────────────────────────────────────────────────────────────

def _stacked_hist(ax, values, is_burst, xlabel, threshold=None, thr_label='', n_bins=30):
    """Non-bursting units in gray at the bottom, bursting units in colour on top."""
    values = np.asarray(values, dtype=float)
    valid  = ~np.isnan(values)
    vals, burst = values[valid], np.asarray(is_burst, dtype=bool)[valid]
    if len(vals) == 0:
        ax.text(0.5, 0.5, 'No valid data', transform=ax.transAxes,
                ha='center', va='center', fontsize=10, color=INK_MUTED)
        ax.set_xlabel(xlabel, color=INK_PRIMARY)
        return
    edges   = np.histogram_bin_edges(vals, bins=n_bins)
    widths  = np.diff(edges)
    centres = edges[:-1] + widths / 2
    n_non, _ = np.histogram(vals[~burst], bins=edges)
    n_bur, _ = np.histogram(vals[burst],  bins=edges)
    if n_non.sum():
        ax.bar(centres, n_non, width=widths, color=NON_BURST_COLOR, edgecolor=SURFACE, linewidth=1.0,
               zorder=3, label=f'Not bursting (n={int(n_non.sum())})')
    if n_bur.sum():
        ax.bar(centres, n_bur, width=widths, bottom=n_non, color=BURST_COLOR, edgecolor=SURFACE,
               linewidth=1.0, zorder=3, label=f'Bursting (n={int(n_bur.sum())})')
    if threshold is not None:
        ax.axvline(threshold, color=MAGMA_3, linewidth=1.5, linestyle='--', zorder=4, label=thr_label)
    mean_v, median_v = float(np.mean(vals)), float(np.median(vals))
    ax.axvline(mean_v,   color=INK_PRIMARY, linewidth=1.5, linestyle='-', zorder=5, label=f'Mean {mean_v:.2f}')
    ax.axvline(median_v, color=INK_PRIMARY, linewidth=1.5, linestyle=':', zorder=5, label=f'Median {median_v:.2f}')
    ax.set_xlabel(xlabel, color=INK_PRIMARY)
    ax.set_ylabel('Number of units', color=INK_PRIMARY)
    _legend_outside(ax)


bursting_flags = df['is_bursting'].astype(bool).values
burst_df       = df[df['is_bursting'].astype(bool)]

fig, axes = plt.subplots(2, 2, figsize=(14, 8), facecolor=SURFACE)
summary_panels = [
    (df['csi'], bursting_flags, 'CSI (%)', 'Complex Spike Index',
     CSI_THRESHOLD, f'CSI = {CSI_THRESHOLD:g}%'),
    (df['burst_power_pct'], bursting_flags,
     f'% of ISIs in [{ISI_BURST_LO_MS:g}–{ISI_BURST_HI_MS:g} ms]', 'Burst power',
     SIGNIFICANCE_PCT, f'{SIGNIFICANCE_PCT:g}%'),
    (burst_df['peak_isi_ms'], np.ones(len(burst_df), dtype=bool), 'Peak burst ISI (ms)',
     'Peak burst ISI (bursting units)', None, ''),
    (burst_df['peak_freq_hz'], np.ones(len(burst_df), dtype=bool), 'Peak burst frequency (Hz)',
     'Peak burst frequency (bursting units)', None, ''),
]
for ax, (vals, flags, xlabel, title, thr, thr_label) in zip(axes.ravel(), summary_panels):
    _style_axes(ax)
    _stacked_hist(ax, pd.to_numeric(vals, errors='coerce').values, flags, xlabel, thr, thr_label)
    ax.set_title(title, color=INK_PRIMARY, fontsize=11, loc='left')
fig.tight_layout()
summary_base = os.path.join(PLOT_DIR, 'BURST_DETECTION_summary')
_save(fig, summary_base)

# ── Excel output ──────────────────────────────────────────────────────────────

summary_rows = [
    ('Units processed',                       len(df)),
    (f'Complex units (CSI > {CSI_THRESHOLD:g}%)', int(df['is_complex'].sum())),
    (f'Burst peak (power >= {SIGNIFICANCE_PCT:g}% and interior KDE peak)', int(df['has_burst_peak'].sum())),
    ('Bursting units',                        int(df['is_bursting'].sum())),
]
for col, label in [('peak_isi_ms', 'Peak ISI (ms)'), ('peak_freq_hz', 'Peak frequency (Hz)'),
                   ('lower_freq_hz', 'Lower half-max frequency (Hz)'),
                   ('upper_freq_hz', 'Upper half-max frequency (Hz)')]:
    v = pd.to_numeric(burst_df[col], errors='coerce').dropna()
    if len(v):
        summary_rows.append((f'{label}, bursting units: mean ± SD (median)',
                             f'{v.mean():.2f} ± {v.std():.2f} ({v.median():.2f})'))
summary_rows += [
    ('', ''),
    ('ROOT_FOLDER', ROOT_FOLDER),
    ('Cluster 0 excluded', EXCLUDE_CLUSTER0),
    ('Burst window (ms)', f'{ISI_BURST_LO_MS:g}–{ISI_BURST_HI_MS:g}'),
    ('ISI ceiling (ms)', ISI_MAX_MS),
    ('KDE bandwidth (ms)', KDE_BW_MS),
    ('Min ISIs', MIN_ISIS),
    ('CSI required', REQUIRE_COMPLEX_SPIKES),
    ('Poisson excess ratio required', 'off' if POISSON_EXCESS_RATIO is None else POISSON_EXCESS_RATIO),
]
summary_df = pd.DataFrame(summary_rows, columns=['Item', 'Value'])

FILL_HEADER = PatternFill(fill_type='solid', fgColor='FF2E4057')
FILL_BURST  = PatternFill(fill_type='solid', fgColor='FF99FF99')
FONT_HEADER = Font(bold=True, color='FFFFFFFF', name='Calibri', size=10)

with pd.ExcelWriter(OUTPUT_EXCEL, engine='openpyxl') as writer:
    df.to_excel(writer, sheet_name='Burst Detection', index=False)
    summary_df.to_excel(writer, sheet_name='Summary', index=False)

    ws = writer.sheets['Burst Detection']
    for cell in ws[1]:
        cell.fill, cell.font = FILL_HEADER, FONT_HEADER
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = 'C2'
    for col_idx, col_name in enumerate(COLUMN_ORDER, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 28 if col_name == 'session' else 15
    burst_col = COLUMN_ORDER.index('is_bursting') + 1
    for row_idx, bursting in enumerate(df['is_bursting'], start=2):
        if bursting:
            ws.cell(row=row_idx, column=burst_col).fill = FILL_BURST

    ss = writer.sheets['Summary']
    for cell in ss[1]:
        cell.fill, cell.font = FILL_HEADER, FONT_HEADER
    ss.column_dimensions['A'].width = 62
    ss.column_dimensions['B'].width = 50
    img = XLImage(f'{summary_base}.png')
    img.width, img.height = img.width * 0.25, img.height * 0.25   # 300 dpi PNG → screen size
    ss.add_image(img, 'D2')

# ── Console summary ───────────────────────────────────────────────────────────

print(f'\nDone. Results saved to {OUTPUT_EXCEL}')
print(f'Plots in {PLOT_DIR}')
for item, value in summary_rows[:summary_rows.index(('', ''))]:
    print(f'  {item}: {value}')
