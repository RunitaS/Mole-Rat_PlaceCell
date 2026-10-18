# -*- coding: utf-8 -*-
"""
Concatenate the per-tetrode L-Ratio / Isolation Distance Excel files written by
Fig1c_lratio_isodist_tetrode_shuffling_v4.py into a single workbook.

For every .ntt file under ROOT_DIR the script looks for its companion
<stem>_LRatio_IsoDist.xlsx (same folder), reads the 'Cluster Quality' sheet,
and adds session / tetrode details parsed from the folder name
(Animal_Arena_DayN_<trial><condition>) plus file-level spike counts read from
the .ntt file size (16 KB header + 304-byte records).

Output workbook (saved in ROOT_DIR) — sheets:
  Summary          overall and per-animal / per-arena counts
  All_Clusters     every cluster from every .ntt file, all metrics
  Per_NTT_File     one row per .ntt file (incl. files with no Excel output)
  High_IsoDist     clusters with Isolation Distance > HIGH_ISO_DIST
  High_LRatio      clusters with L-Ratio > HIGH_L_RATIO
  Thresholds       values / definitions used

z-scores are recomputed from the rounded null mean ± SD stored in the Excel
files (3 decimals), so they are approximate.
"""

import os
import re
import numpy as np # type: ignore
import pandas as pd # type: ignore
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side # type: ignore
from openpyxl.utils import get_column_letter # type: ignore

# ── Configuration ──────────────────────────────────────────────────────────────

ROOT_DIR = r'X:\NMR_group_data\Runita\Analysis\Thesis\Corr_Data_SpkQltyFilt\SpikeQualityFilt_v2\QualFilt_DataConcat'

XLSX_SUFFIX  = '_LRatio_IsoDist.xlsx'                 # written by the Fig1c v4 code
OUT_NAME     = 'AllNtt_LRatioIsoDist_Concatenated.xlsx'

# Suspicious-value cutoffs
HIGH_ISO_DIST = 100.0   # Isolation Distance (d²) above this is suspicious
HIGH_L_RATIO  = 4.0     # L-Ratio above this is suspicious

# Quality thresholds used by the Fig1c code (for reference in the Thresholds sheet)
L_RATIO_THRESHOLD  = 0.2
ISO_DIST_THRESHOLD = 10.0
SIG_ALPHA          = 0.05
N_FEATURES         = 8      # FD dims in Fig1c (Energy + WavePC1 × 4 channels)

# ── Neuralynx .ntt layout (for spike counts from file size) ──────────────────

NTT_HEADER_BYTES = 16 * 1024
NTT_RECORD_BYTES = 8 + 4 + 4 + 8 * 4 + 32 * 4 * 2   # = 304

# ── Excel styling ──────────────────────────────────────────────────────────────

FILL_HEADER = PatternFill(fill_type='solid', fgColor='FF2E4057')
FILL_GREEN  = PatternFill(fill_type='solid', fgColor='FF90EE90')
FILL_RED    = PatternFill(fill_type='solid', fgColor='FFF4A6A6')
FILL_ORANGE = PatternFill(fill_type='solid', fgColor='FFFFD08A')
FONT_HEADER = Font(bold=True, color='FFFFFFFF', name='Calibri', size=10)
FONT_BODY   = Font(name='Calibri', size=10)
THIN_BORDER = Border(left=Side(style='thin'), right=Side(style='thin'),
                     top=Side(style='thin'), bottom=Side(style='thin'))

# ── Source-column mapping (matched by header prefix, robust to threshold text) ─

SOURCE_COLS = [
    ('Cluster ID',         'Cluster'),
    ('N Spikes',           'N_Spikes'),
    ('Peak Amp',           'PeakAmp_uV'),
    ('L-Ratio',            'L_Ratio'),
    ('Isolation Distance', 'IsoDist'),
    ('Quality',            'Quality'),
    ('NN Shuffle',         '_shuffle_valid_pool'),
    ('Null L-Ratio',       '_null_lratio'),
    ('p(L-Ratio)',         'p_LRatio'),
    ('Null IsoDist',       '_null_isodist'),
    ('p(IsoDist)',         'p_IsoDist'),
    ('Shuffle Sig',        'Shuffle_Sig'),
]

SESSION_RE = re.compile(r'^(?P<animal>[^_]+)_(?P<arena>[^_]+)_Day(?P<day>\d+)_?'
                        r'(?P<trial>\d+)?_?(?P<cond>.*)$')


# ── Parsing helpers ───────────────────────────────────────────────────────────

def parse_session(folder_name: str) -> dict:
    m = SESSION_RE.match(folder_name)
    if not m:
        return {'Animal': '', 'Arena': '', 'Day': np.nan, 'Trial': np.nan, 'Condition': ''}
    return {
        'Animal':    m['animal'],
        'Arena':     m['arena'],
        'Day':       int(m['day']),
        'Trial':     int(m['trial']) if m['trial'] else np.nan,
        'Condition': m['cond'],
    }


def ntt_spike_count(path: str) -> float:
    try:
        n_bytes = os.path.getsize(path) - NTT_HEADER_BYTES
    except OSError:
        return np.nan
    return n_bytes // NTT_RECORD_BYTES if n_bytes >= 0 else np.nan


def _split_pair(series: pd.Series, sep: str):
    """'0.508 ± 0.015' → (0.508, 0.015);  '1000 / 1650' → (1000, 1650)."""
    num   = r'\s*([-+0-9.eE]+)\s*'
    parts = series.astype(str).str.extract('^' + num + re.escape(sep) + num + '$')
    return (pd.to_numeric(parts[0], errors='coerce'),
            pd.to_numeric(parts[1], errors='coerce'))


def read_cluster_sheet(xlsx_path: str) -> pd.DataFrame:
    """Cluster rows of one Fig1c output file (method footnotes dropped)."""
    raw = pd.read_excel(xlsx_path, sheet_name=0)

    rename = {}
    for col in raw.columns:
        for prefix, name in SOURCE_COLS:
            if str(col).startswith(prefix) and name not in rename.values():
                rename[col] = name
                break
    df = raw.rename(columns=rename)

    df['Cluster'] = pd.to_numeric(df['Cluster'], errors='coerce')
    df = df[df['Cluster'].notna()].copy()
    df['Cluster'] = df['Cluster'].astype(int)

    for c in ['N_Spikes', 'PeakAmp_uV', 'L_Ratio', 'IsoDist', 'p_LRatio', 'p_IsoDist']:
        df[c] = pd.to_numeric(df[c], errors='coerce')

    df['Shuffle_N_Valid'], df['Shuffle_Pool'] = _split_pair(df['_shuffle_valid_pool'], '/')
    df['Null_LRatio_Mean'], df['Null_LRatio_SD'] = _split_pair(df['_null_lratio'], '±')
    df['Null_IsoDist_Mean'], df['Null_IsoDist_SD'] = _split_pair(df['_null_isodist'], '±')
    df['Shuffle_Sig'] = df['Shuffle_Sig'].fillna('N/A')

    # positive z = real cluster better separated than the NN-shuffle null
    df['z_LRatio'] = np.where(df['Null_LRatio_SD'] > 0,
                              (df['Null_LRatio_Mean'] - df['L_Ratio']) / df['Null_LRatio_SD'], np.nan)
    df['z_IsoDist'] = np.where(df['Null_IsoDist_SD'] > 0,
                               (df['IsoDist'] - df['Null_IsoDist_Mean']) / df['Null_IsoDist_SD'], np.nan)
    return df.drop(columns=[c for c in df.columns if c.startswith('_')])


# ── Collection ────────────────────────────────────────────────────────────────

def collect(root_dir: str):
    cluster_frames, file_rows = [], []

    for dirpath, _dirs, files in os.walk(root_dir):
        for fname in sorted(files):
            if not fname.lower().endswith('.ntt'):
                continue
            ntt_path  = os.path.join(dirpath, fname)
            stem      = os.path.splitext(fname)[0]
            xlsx_path = os.path.join(dirpath, stem + XLSX_SUFFIX)
            session   = os.path.basename(dirpath)

            info = {
                'Session':   session,
                **parse_session(session),
                'Tetrode':   stem,
                'NTT_File':  fname,
                'RelFolder': os.path.relpath(dirpath, root_dir),
            }
            total = ntt_spike_count(ntt_path)
            frow  = {**info, 'Total_Spikes_File': total, 'NTT_Path': ntt_path,
                     'Excel_Path': xlsx_path if os.path.isfile(xlsx_path) else ''}

            if not os.path.isfile(xlsx_path):
                frow['Status'] = 'No Excel output (no sorted clusters or not processed)'
                file_rows.append(frow)
                continue
            try:
                df = read_cluster_sheet(xlsx_path)
            except Exception as exc:
                frow['Status'] = f'Excel read error: {exc}'
                file_rows.append(frow)
                print(f'  ERROR {xlsx_path}: {exc}')
                continue

            n_sorted = df['N_Spikes'].sum()
            df['Total_Spikes_File']   = total
            df['Unsorted_Spikes_C0']  = total - n_sorted
            df['Background_Spikes']   = total - df['N_Spikes']
            df['Cluster_Frac_File']   = df['N_Spikes'] / total
            df['N_Clusters_Tetrode']  = len(df)
            df['IsoDist_Undefined']   = np.where(df['N_Spikes'] > total / 2,
                                                 'Yes (cluster > 50% of spikes)', '')
            # Covariance of N_FEATURES dims needs > N_FEATURES spikes; below that
            # mahal() inverts a near-singular matrix and d² explodes
            df['Cov_Singular'] = np.where(df['N_Spikes'] <= N_FEATURES,
                                          f'Yes (N spikes <= {N_FEATURES} features)', '')
            df['High_IsoDist'] = df['IsoDist'] > HIGH_ISO_DIST
            df['High_LRatio']  = df['L_Ratio'] > HIGH_L_RATIO
            for k, v in reversed(list(info.items())):
                df.insert(0, k, v)
            df['NTT_Path']   = ntt_path
            df['Excel_Path'] = xlsx_path
            cluster_frames.append(df)

            frow.update({
                'Status':               'OK',
                'Unsorted_Spikes_C0':   total - n_sorted,
                'N_Clusters':           len(df),
                'N_Good_Threshold':     int((df['Quality'] == 'Good').sum()),
                'N_Shuffle_Sig':        int((df['Shuffle_Sig'] == 'Yes').sum()),
                'N_IsoDist_NaN':        int(df['IsoDist'].isna().sum()),
                f'N_IsoDist_>{HIGH_ISO_DIST:g}': int(df['High_IsoDist'].sum()),
                f'N_LRatio_>{HIGH_L_RATIO:g}':   int(df['High_LRatio'].sum()),
                'Median_LRatio':        df['L_Ratio'].median(),
                'Median_IsoDist':       df['IsoDist'].median(),
                'Max_IsoDist':          df['IsoDist'].max(),
                'Max_LRatio':           df['L_Ratio'].max(),
            })
            file_rows.append(frow)

    clusters = pd.concat(cluster_frames, ignore_index=True) if cluster_frames else pd.DataFrame()
    files_df = pd.DataFrame(file_rows)
    # keep Status / paths at the end of the per-file sheet
    tail = ['Status', 'NTT_Path', 'Excel_Path']
    files_df = files_df[[c for c in files_df.columns if c not in tail] + tail]
    return clusters, files_df


# ── Summary sheet ─────────────────────────────────────────────────────────────

def _group_summary(df: pd.DataFrame, by) -> pd.DataFrame:
    g = df.groupby(by, dropna=False)
    out = pd.DataFrame({
        'N_NTT_Files':          g['NTT_Path'].nunique(),
        'N_Clusters':           g.size(),
        'N_Good_Threshold':     g['Quality'].apply(lambda s: (s == 'Good').sum()),
        'N_Shuffle_Sig':        g['Shuffle_Sig'].apply(lambda s: (s == 'Yes').sum()),
        'N_IsoDist_NaN':        g['IsoDist'].apply(lambda s: s.isna().sum()),
        f'N_IsoDist_>{HIGH_ISO_DIST:g}': g['High_IsoDist'].sum(),
        f'N_LRatio_>{HIGH_L_RATIO:g}':   g['High_LRatio'].sum(),
        'Median_LRatio':        g['L_Ratio'].median(),
        'Median_IsoDist':       g['IsoDist'].median(),
    }).reset_index()
    return out


def build_summary(clusters: pd.DataFrame, files_df: pd.DataFrame):
    overall = pd.DataFrame([
        ('.ntt files found',                       len(files_df)),
        ('.ntt files with Excel output',           int((files_df['Status'] == 'OK').sum())),
        ('.ntt files without Excel output',        int((files_df['Status'] != 'OK').sum())),
        ('Total clusters',                         len(clusters)),
        (f'Good (L-Ratio < {L_RATIO_THRESHOLD} & IsoDist > {ISO_DIST_THRESHOLD:g})',
                                                   int((clusters['Quality'] == 'Good').sum())),
        (f'NN-shuffle significant (p < {SIG_ALPHA} both)',
                                                   int((clusters['Shuffle_Sig'] == 'Yes').sum())),
        ('Isolation Distance undefined (NaN)',     int(clusters['IsoDist'].isna().sum())),
        (f'Isolation Distance > {HIGH_ISO_DIST:g} (suspicious)', int(clusters['High_IsoDist'].sum())),
        (f'L-Ratio > {HIGH_L_RATIO:g} (suspicious)',             int(clusters['High_LRatio'].sum())),
        ('Both suspicious flags',                  int((clusters['High_IsoDist'] & clusters['High_LRatio']).sum())),
        (f'Singular covariance (N spikes <= {N_FEATURES})', int((clusters['Cov_Singular'] != '').sum())),
        (f'  of which IsoDist > {HIGH_ISO_DIST:g}',
                                                   int(((clusters['Cov_Singular'] != '') & clusters['High_IsoDist']).sum())),
        ('Median L-Ratio',                         round(float(clusters['L_Ratio'].median()), 4)),
        ('Median Isolation Distance',              round(float(clusters['IsoDist'].median()), 4)),
    ], columns=['Metric', 'Value'], dtype=object)
    return overall, _group_summary(clusters, ['Animal']), _group_summary(clusters, ['Animal', 'Arena'])


# ── Excel writing / styling ───────────────────────────────────────────────────

def _style_sheet(ws, df: pd.DataFrame, start_row: int = 1, row_fill=None):
    """Header style, borders, widths, freeze panes, autofilter."""
    n_rows, n_cols = df.shape
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=start_row, column=c)
        cell.fill, cell.font, cell.border = FILL_HEADER, FONT_HEADER, THIN_BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[start_row].height = 30

    for r in range(start_row + 1, start_row + n_rows + 1):
        fill = row_fill(r - start_row - 1) if row_fill else None
        for c in range(1, n_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.border = FONT_BODY, THIN_BORDER
            if fill is not None:
                cell.fill = fill

    for c, col in enumerate(df.columns, 1):
        sample = [str(col)] + [str(v) for v in df[col].head(200).tolist()]
        width  = min(max(len(s) for s in sample) + 2, 60)
        ws.column_dimensions[get_column_letter(c)].width = max(width, 8)

    if start_row == 1 and n_rows > 0:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(n_cols)}{n_rows + 1}'


def _cluster_fill(df: pd.DataFrame):
    hi_iso, hi_lr, good = df['High_IsoDist'].values, df['High_LRatio'].values, \
                          (df['Quality'] == 'Good').values
    def f(i):
        if hi_iso[i] or hi_lr[i]:
            return FILL_RED if hi_iso[i] and hi_lr[i] else FILL_ORANGE
        return FILL_GREEN if good[i] else None
    return f


def _round(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for c in out.select_dtypes(include='float').columns:
        out[c] = out[c].round(4)
    return out


def write_workbook(clusters, files_df, out_path):
    overall, by_animal, by_arena = build_summary(clusters, files_df)

    sort_cols = ['Animal', 'Arena', 'Day', 'Trial', 'Session', 'Tetrode', 'Cluster']
    clusters  = clusters.sort_values(sort_cols, kind='stable').reset_index(drop=True)
    files_df  = files_df.sort_values(['Animal', 'Arena', 'Day', 'Trial', 'Session', 'Tetrode'],
                                     kind='stable').reset_index(drop=True)

    high_iso = clusters[clusters['High_IsoDist']].sort_values('IsoDist', ascending=False)
    high_lr  = clusters[clusters['High_LRatio']].sort_values('L_Ratio', ascending=False)

    thresholds = pd.DataFrame([
        ('ROOT_DIR', ROOT_DIR),
        ('Suspicious Isolation Distance', f'> {HIGH_ISO_DIST:g} (d²)'),
        ('Suspicious L-Ratio',            f'> {HIGH_L_RATIO:g}'),
        ('Good (Fig1c)',                  f'L-Ratio < {L_RATIO_THRESHOLD} AND IsoDist > {ISO_DIST_THRESHOLD:g}'),
        ('Shuffle Sig (Fig1c)',           f'p(L-Ratio) < {SIG_ALPHA} AND p(IsoDist) < {SIG_ALPHA}'),
        ('Total_Spikes_File',             f'(file size - {NTT_HEADER_BYTES}) / {NTT_RECORD_BYTES} bytes per record'),
        ('Unsorted_Spikes_C0',            'Total_Spikes_File - sum of sorted-cluster spikes (cell 0)'),
        ('Background_Spikes',             'Total_Spikes_File - cluster spikes (spikes IsoDist / L-Ratio are computed against)'),
        ('IsoDist_Undefined',             'Cluster holds > 50% of the file spikes, so IsoDist is NaN (L-Ratio still computed)'),
        ('Cov_Singular',                  f'N spikes <= {N_FEATURES} feature dims: cluster covariance is singular, '
                                          'IsoDist / L-Ratio are numerical artefacts'),
        ('z_LRatio / z_IsoDist',          'From rounded null mean ± SD in the Excel files; positive = better than NN-shuffle null'),
        ('Row colours',                   'red = both suspicious, orange = one suspicious, green = Good'),
    ], columns=['Item', 'Definition'])

    with pd.ExcelWriter(out_path, engine='openpyxl') as xw:
        # Summary: three stacked tables
        r = 0
        blocks = [('Overall', overall), ('Per animal', by_animal), ('Per animal × arena', by_arena)]
        positions = []
        for title, tbl in blocks:
            pd.DataFrame([[title]]).to_excel(xw, sheet_name='Summary', index=False,
                                             header=False, startrow=r)
            _round(tbl).to_excel(xw, sheet_name='Summary', index=False, startrow=r + 1)
            positions.append((r + 2, tbl))
            r += len(tbl) + 4
        ws = xw.sheets['Summary']
        for hdr_row, tbl in positions:
            ws.cell(row=hdr_row - 1, column=1).font = Font(bold=True, size=12, name='Calibri')
            _style_sheet(ws, tbl, start_row=hdr_row)

        for name, df, fill in [
            ('All_Clusters', clusters, _cluster_fill(clusters)),
            ('Per_NTT_File', files_df, None),
            ('High_IsoDist', high_iso, None),
            ('High_LRatio',  high_lr,  None),
            ('Thresholds',   thresholds, None),
        ]:
            df = _round(df).reset_index(drop=True)
            df.to_excel(xw, sheet_name=name, index=False)
            _style_sheet(xw.sheets[name], df, row_fill=fill)

    return overall, high_iso, high_lr


# ── Main ──────────────────────────────────────────────────────────────────────

def main(root_dir: str = ROOT_DIR):
    print(f'Scanning: {root_dir}')
    clusters, files_df = collect(root_dir)
    if clusters.empty:
        print('No cluster rows found — nothing written.')
        return

    out_path = os.path.join(root_dir, OUT_NAME)
    overall, high_iso, high_lr = write_workbook(clusters, files_df, out_path)

    print()
    for _, row in overall.iterrows():
        print(f'  {row["Metric"]:<55} {row["Value"]}')
    print(f'\nSaved: {out_path}')


if __name__ == '__main__':
    main()
