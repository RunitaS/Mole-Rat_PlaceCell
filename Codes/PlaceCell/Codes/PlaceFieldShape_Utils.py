# -*- coding: utf-8 -*-
"""
Place-field shape from the bins of each field (the 'bin_coords' column of the
'PlaceFields' sheet written by PlaceCellChar_FieldDetect_Main_v3.py, e.g.
'0-2;0-3;1-2' = bins (x=0, y=2), (x=0, y=3), (x=1, y=2)).

Each 2-D field is measured with calipers on the centres of its bins. The width
of a field along an axis is the distance between the centres of its two
outermost bins along that axis plus one bin, so a straight row of n bins is
n bins long and 1 bin wide:

  - max width   : widest extent over all axes (the axis with the most bins);
                  the distance between the two farthest bin centres + 1 bin
  - min width   : narrowest extent over all axes (the axis with the fewest
                  bins); for a convex outline this is always measured across
                  one of its edges, so it is found on the convex hull of the
                  bin centres (a single row of bins: 1 bin)
  - axis angles : orientation of the max- and min-width axes (deg, 0-180,
                  counter-clockwise from the +x bin axis), and the acute angle
                  between them (0-90). The two axes need not be perpendicular,
                  e.g. a 3 x 3 square field: max width along the diagonal, min
                  width along a side, 45 deg apart.
  - width ratio : max width / min width (1 = round/square, > 1 = elongated)
  - mean width  : (max width + min width) / 2

Widths are given in bins and in cm (bins x bin side, bin side =
sqrt(area_cm2 / n_bins)).

Circular-track fields ('bin_coords' are 1-D angular bins, e.g. '25;26;27')
have no 2-D shape; their shape columns are left blank (NaN).
"""

from copy import copy

import numpy as np
import pandas as pd
import openpyxl
from scipy.spatial import ConvexHull

# Columns added to the 'PlaceFields' sheet, in order.
SHAPE_COLUMNS = [
    'fld_max_width_bins', 'fld_min_width_bins',
    'fld_max_width_cm',   'fld_min_width_cm',
    'fld_max_axis_deg',   'fld_min_axis_deg',   'fld_axis_angle_deg',
    'fld_width_ratio',
    'fld_mean_width_bins', 'fld_mean_width_cm',
]


def _parse_bins(bin_coords):
    """(n, 2) float array of 'x-y' bin indices, or None if `bin_coords` is
    blank or not 2-D (circular-track fields list single angular bins)."""
    if not isinstance(bin_coords, str) or '-' not in bin_coords:
        return None
    return np.array([[float(v) for v in tok.split('-')]
                     for tok in bin_coords.split(';') if tok.strip()])


def _axis_deg(v) -> float:
    return float(np.degrees(np.arctan2(v[1], v[0])) % 180.0)


def field_shape(bin_coords) -> dict:
    """Shape of one field (widths in bins, see module docstring). All values
    NaN if the field has no 2-D bins."""
    out = dict.fromkeys(['max_width', 'min_width', 'max_axis_deg', 'min_axis_deg',
                         'axis_angle_deg'], np.nan)
    pts = _parse_bins(bin_coords)
    if pts is None or not len(pts):
        return out
    pts = np.unique(pts, axis=0)
    if len(pts) == 1:
        out.update(max_width=1.0, min_width=1.0)
        return out

    # Max width: the two farthest bin centres.
    d = pts[:, None, :] - pts[None, :, :]
    dist = np.hypot(d[..., 0], d[..., 1])
    i, j = np.unravel_index(np.argmax(dist), dist.shape)
    max_dir = d[i, j]
    out['max_width'] = dist[i, j] + 1.0
    out['max_axis_deg'] = _axis_deg(max_dir)

    # Min width: narrowest projection, measured across each convex-hull edge.
    centred = pts - pts.mean(axis=0)
    if np.linalg.matrix_rank(centred, tol=1e-9) < 2:   # all bins in one line
        min_w, min_dir = 0.0, np.array([-max_dir[1], max_dir[0]])
    else:
        hull = pts[ConvexHull(pts).vertices]
        edges = np.roll(hull, -1, axis=0) - hull
        normals = np.column_stack([-edges[:, 1], edges[:, 0]])
        normals /= np.hypot(normals[:, 0], normals[:, 1])[:, None]
        proj = hull @ normals.T                  # (n_vertices, n_edges)
        widths = proj.max(axis=0) - proj.min(axis=0)
        k = int(np.argmin(widths))
        min_w, min_dir = widths[k], normals[k]
    out['min_width'] = min_w + 1.0
    out['min_axis_deg'] = _axis_deg(min_dir)

    between = abs(out['max_axis_deg'] - out['min_axis_deg']) % 180.0
    out['axis_angle_deg'] = min(between, 180.0 - between)
    return out


def add_field_shape(fields: pd.DataFrame) -> pd.DataFrame:
    """Copy of the 'PlaceFields' dataframe with the SHAPE_COLUMNS added."""
    fields = fields.copy()
    shp = pd.DataFrame([field_shape(b) for b in fields['bin_coords']], index=fields.index)
    bin_cm = np.sqrt(pd.to_numeric(fields['area_cm2'], errors='coerce')
                     / pd.to_numeric(fields['n_bins'], errors='coerce'))

    fields['fld_max_width_bins']  = shp['max_width']
    fields['fld_min_width_bins']  = shp['min_width']
    fields['fld_max_width_cm']    = shp['max_width'] * bin_cm
    fields['fld_min_width_cm']    = shp['min_width'] * bin_cm
    fields['fld_max_axis_deg']    = shp['max_axis_deg']
    fields['fld_min_axis_deg']    = shp['min_axis_deg']
    fields['fld_axis_angle_deg']  = shp['axis_angle_deg']
    fields['fld_width_ratio']     = shp['max_width'] / shp['min_width']
    fields['fld_mean_width_bins'] = (shp['max_width'] + shp['min_width']) / 2
    fields['fld_mean_width_cm']   = fields['fld_mean_width_bins'] * bin_cm
    return fields


def save_shape_columns(in_path: str, out_path: str, fields: pd.DataFrame,
                       sheet: str = 'PlaceFields'):
    """Write a copy of workbook `in_path` to `out_path` with the SHAPE_COLUMNS
    of `fields` (same rows, same order as `sheet`) written into `sheet` --
    into the first empty columns after the existing ones, or over columns of
    the same name if they already exist. Other sheets are copied unchanged.
    `in_path` itself is not modified."""
    wb = openpyxl.load_workbook(in_path)
    ws = wb[sheet]
    if ws.max_row - 1 != len(fields):
        raise ValueError(f"'{sheet}' has {ws.max_row - 1} data rows, "
                         f'but {len(fields)} field rows were computed')

    header = [c.value for c in ws[1]]
    while header and header[-1] is None:
        header.pop()
    col_of = {name: i + 1 for i, name in enumerate(header) if name is not None}
    next_col = len(header) + 1
    for name in SHAPE_COLUMNS:
        if name not in col_of:
            col_of[name] = next_col
            next_col += 1
        col = col_of[name]
        head = ws.cell(row=1, column=col, value=name)
        head.font = copy(ws.cell(row=1, column=1).font)
        for r, v in enumerate(fields[name].to_numpy(dtype=float), start=2):
            ws.cell(row=r, column=col, value=None if np.isnan(v) else round(float(v), 4))
    wb.save(out_path)
